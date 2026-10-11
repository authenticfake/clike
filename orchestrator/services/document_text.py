"""Plain-text extraction of office and PDF documents (PDF, DOCX, XLSX/XLS/XLSB, PPTX).

Shared by RAG indexing (routes/rag.py) and the project context sources (services/context_sources.py).
"""
from __future__ import annotations

import io
import logging
import os

import docx
import openpyxl  # .xlsx
import xlrd  # .xls (legacy)
from pdfminer.high_level import extract_text
from pptx import Presentation  # .pptx
from pyxlsb import open_workbook as open_xlsb  # .xlsb

log = logging.getLogger(__name__)


def _extract_text_from_pdf_bytes(raw: bytes) -> str:
    # pdfminer.six
    try:
        return extract_text(io.BytesIO(raw)) or ""
    except Exception as e:
        log.warning("PDF extract failed: %s", e)
        return ""


def _extract_text_from_docx_bytes(raw: bytes) -> str:
    # python-docx
    try:
        doc = docx.Document(io.BytesIO(raw))
        parts = []
        # paragraphs
        for p in doc.paragraphs:
            t = (p.text or "").strip()
            if t:
                parts.append(t)
        # tables (cells)
        for tbl in getattr(doc, "tables", []):
            for row in tbl.rows:
                for cell in row.cells:
                    t = (cell.text or "").strip()
                    if t:
                        parts.append(t)
        return "\n".join(parts)
    except Exception as e:
        log.warning("DOCX extract failed: %s", e)
        return ""


def _extract_text_from_xlsx_bytes(raw: bytes) -> str:
    if not openpyxl:
        log.warning("openpyxl non disponibile: skip xlsx")
        return ""
    try:
        wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
        parts = []
        for ws in wb.worksheets:
            parts.append(f"# Sheet: {ws.title}")
            for row in ws.iter_rows(values_only=True):
                cells = [str(c) if c is not None else "" for c in row]
                line = "\t".join(cells).strip()
                if line:
                    parts.append(line)
        return "\n".join(parts)
    except Exception as e:
        log.warning("XLSX extract failed: %s", e)
        return ""


def _extract_text_from_xls_bytes(raw: bytes) -> str:
    # Requires xlrd>=2.0 (reads .xls only)
    if not xlrd:
        log.warning("xlrd non disponibile: skip xls")
        return ""
    try:
        book = xlrd.open_workbook(file_contents=raw)
        parts = []
        for si in range(book.nsheets):
            sh = book.sheet_by_index(si)
            parts.append(f"# Sheet: {sh.name}")
            for r in range(sh.nrows):
                row = [str(sh.cell_value(r, c)) for c in range(sh.ncols)]
                line = "\t".join(row).strip()
                if line:
                    parts.append(line)
        return "\n".join(parts)
    except Exception as e:
        log.warning("XLS extract failed: %s", e)
        return ""


def _extract_text_from_xlsb_bytes(raw: bytes) -> str:
    if not open_xlsb:
        log.warning("pyxlsb non disponibile: skip xlsb")
        return ""
    try:
        parts = []
        with open_xlsb(io.BytesIO(raw)) as wb:
            for sheet_name in wb.sheets:
                parts.append(f"# Sheet: {sheet_name}")
                with wb.get_sheet(sheet_name) as sh:
                    for row in sh.rows():
                        vals = [str(c.v) if c.v is not None else "" for c in row]
                        line = "\t".join(vals).strip()
                        if line:
                            parts.append(line)
        return "\n".join(parts)
    except Exception as e:
        log.warning("XLSB extract failed: %s", e)
        return ""


def _extract_text_from_pptx_bytes(raw: bytes) -> str:
    if not Presentation:
        log.warning("python-pptx non disponibile: skip pptx")
        return ""
    try:
        prs = Presentation(io.BytesIO(raw))
        parts = []
        for i, slide in enumerate(prs.slides, start=1):
            parts.append(f"# Slide {i}")
            for shape in slide.shapes:
                if hasattr(shape, "text_frame") and shape.text_frame:
                    for para in shape.text_frame.paragraphs:
                        text = "".join(run.text or "" for run in para.runs).strip()
                        if text:
                            parts.append(text)
                elif hasattr(shape, "text") and shape.text:
                    t = (shape.text or "").strip()
                    if t:
                        parts.append(t)
        return "\n".join(parts)
    except Exception as e:
        log.warning("PPTX extract failed: %s", e)
        return ""


_EXTRACTORS = {
    ".pdf": _extract_text_from_pdf_bytes,
    ".docx": _extract_text_from_docx_bytes,
    ".xlsx": _extract_text_from_xlsx_bytes,
    ".xls": _extract_text_from_xls_bytes,
    ".xlsb": _extract_text_from_xlsb_bytes,
    ".pptx": _extract_text_from_pptx_bytes,
}

CONVERTIBLE_EXTENSIONS = frozenset(_EXTRACTORS)


def extract_document_text(path: str, raw: bytes) -> str:
    """Text of an office/PDF document ("" when the format is unknown or extraction fails)."""
    extractor = _EXTRACTORS.get(os.path.splitext(str(path or ""))[1].lower())
    return extractor(raw) if extractor else ""
