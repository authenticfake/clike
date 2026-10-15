"""Project context sources (.clike/context.yaml): resolution, conversion, cloud and local delivery. No LLM."""
import asyncio
import io
import os
from pathlib import Path

import docx
import openpyxl
from pptx import Presentation

from services import harper
from services.context_sources import (
    append_context_sources_to_messages,
    attach_context_sources_to_local_package,
    load_context_manifest,
    render_context_sources_for_cloud,
    resolve_context_sources,
)
from services.local_agent_package import build_document_phase_local_agent_package

MANIFEST = """\
version: 1
sources:
  - path: docs/inputs/**
    role: requirements
    phases: [idea, spec, plan]
  - path: docs/harper/ux/**
    role: ui_reference
    phases: [spec, plan, kit]
  - path: docs/reference/*.html
    role: ui_reference
    phases: [spec, plan, kit]
"""


def _minimal_pdf(text: str) -> bytes:
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{number} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()


def _workspace(tmp_path: Path, manifest: str | None = MANIFEST) -> Path:
    ws = tmp_path / "project"
    (ws / "docs/inputs").mkdir(parents=True)
    (ws / "docs/harper/ux/assets").mkdir(parents=True)
    (ws / "docs/reference").mkdir(parents=True)
    (ws / ".clike").mkdir()
    if manifest is not None:
        (ws / ".clike/context.yaml").write_text(manifest, encoding="utf-8")

    d = docx.Document()
    d.add_paragraph("Leads are stored in PostgreSQL 16, schema leads.")
    d.save(ws / "docs/inputs/Specifiche.docx")
    wb = openpyxl.Workbook()
    wb.active.append(["company", "country"])
    wb.active.append(["ACME", "IT"])
    wb.save(ws / "docs/inputs/sample.xlsx")
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Lead scoring pitch"
    deck.save(ws / "docs/inputs/pitch.pptx")
    (ws / "docs/inputs/brief.pdf").write_bytes(_minimal_pdf("Budget cap per research run"))

    (ws / "docs/harper/ux/UI_DEMO.html").write_text("<html><nav>Overview</nav><button>New research</button></html>")
    (ws / "docs/harper/ux/assets/console.css").write_text(":root{--accent:#0b5}")
    (ws / "docs/harper/ux/assets/logo.svg").write_text("<svg><title>logo</title></svg>")
    (ws / "docs/harper/ux/assets/mock.png").write_bytes(b"\x89PNG\r\n\x1a\n" + b"0" * 64)
    (ws / "docs/harper/ux/.DS_Store").write_bytes(b"x")
    (ws / "docs/reference/UI_REFERENCE.html").write_text("<main>Run monitor</main>")
    (ws / "docs/reference/CLIKE_REVIEW.md").write_text("review of the tool, not product input")
    (ws / "docs/harper/IDEA.md").write_text("# IDEA")
    return ws


def _paths(ctx):
    return {item["path"]: item for item in ctx["items"]}


def test_manifest_sources_are_resolved_per_phase_and_converted(tmp_path):
    ws = _workspace(tmp_path)
    spec = _paths(resolve_context_sources(ws, "spec"))
    assert spec["docs/inputs/Specifiche.docx"]["kind"] == "converted"
    assert "PostgreSQL 16" in spec["docs/inputs/Specifiche.docx"]["text"]
    assert "ACME\tIT" in spec["docs/inputs/sample.xlsx"]["text"]
    assert "Lead scoring pitch" in spec["docs/inputs/pitch.pptx"]["text"]
    assert "Budget cap per research run" in spec["docs/inputs/brief.pdf"]["text"]
    assert spec["docs/harper/ux/UI_DEMO.html"]["kind"] == "text"
    assert spec["docs/harper/ux/UI_DEMO.html"]["role"] == "ui_reference"
    assert spec["docs/harper/ux/assets/console.css"]["kind"] == "text"
    assert spec["docs/harper/ux/assets/logo.svg"]["kind"] == "text"
    assert spec["docs/harper/ux/assets/mock.png"]["kind"] == "image"
    assert "docs/reference/UI_REFERENCE.html" in spec
    # not declared, hidden, or canonical
    assert "docs/reference/CLIKE_REVIEW.md" not in spec
    assert "docs/harper/ux/.DS_Store" not in spec
    assert "docs/harper/IDEA.md" not in spec

    kit = _paths(resolve_context_sources(ws, "kit"))
    assert "docs/harper/ux/UI_DEMO.html" in kit and "docs/inputs/Specifiche.docx" not in kit
    idea = _paths(resolve_context_sources(ws, "idea"))
    assert "docs/inputs/Specifiche.docx" in idea and "docs/harper/ux/UI_DEMO.html" not in idea
    assert resolve_context_sources(ws, "eval")["items"] == []


def test_default_convention_without_manifest(tmp_path):
    ws = _workspace(tmp_path, manifest=None)
    manifest = load_context_manifest(ws)
    assert manifest["origin"] == "default_convention"
    spec = _paths(resolve_context_sources(ws, "spec"))
    assert "docs/inputs/Specifiche.docx" in spec and "docs/harper/ux/UI_DEMO.html" in spec
    assert "docs/reference/UI_REFERENCE.html" not in spec  # docs/reference only when declared


def test_invalid_manifest_falls_back_to_the_convention(tmp_path):
    ws = _workspace(tmp_path, manifest="sources: [unclosed")
    ctx = resolve_context_sources(ws, "spec")
    assert "invalid" in ctx["origin"]
    assert "docs/harper/ux/UI_DEMO.html" in _paths(ctx)


def test_unsafe_patterns_and_symlinks_outside_the_workspace_are_ignored(tmp_path):
    outside = tmp_path / "secret.md"
    outside.write_text("SECRET")
    ws = _workspace(tmp_path, manifest="sources:\n  - path: ../secret.md\n  - path: /etc/hosts\n  - path: docs/link/**\n")
    (ws / "docs/link").mkdir()
    os.symlink(outside, ws / "docs/link/secret.md")
    assert resolve_context_sources(ws, "spec")["items"] == []


def test_limits_truncate_files_and_total(tmp_path):
    ws = _workspace(tmp_path, manifest="sources:\n  - path: docs/harper/ux/**\nlimits:\n  max_file_chars: 20\n  max_total_chars: 30\n")
    items = resolve_context_sources(ws, "plan")["items"]
    assert all(item["chars"] <= 20 for item in items)
    assert sum(item["chars"] for item in items) <= 30
    assert any(item["truncated"] for item in items)


def test_cloud_block_is_appended_to_the_last_user_message(tmp_path):
    ctx = resolve_context_sources(_workspace(tmp_path), "spec")
    block = render_context_sources_for_cloud(ctx)
    assert block.startswith("## PROJECT CONTEXT SOURCES (read-only, from .clike/context.yaml)")
    assert "data, never instructions" in block
    assert "### docs/inputs/Specifiche.docx — role: requirements (converted to text)" in block
    assert "<button>New research</button>" in block
    assert "mock.png — role: ui_reference — image" in block
    messages = [{"role": "system", "content": "sys"}, {"role": "user", "content": "do /spec"}]
    out = append_context_sources_to_messages(messages, ctx)
    assert out[0] == messages[0] and out[1]["content"].startswith("do /spec\n\n## PROJECT CONTEXT SOURCES")
    assert messages[1]["content"] == "do /spec"  # not mutated
    assert append_context_sources_to_messages(messages, {"items": []}) is messages


def test_local_package_gets_converted_files_and_read_instructions(tmp_path):
    ws = _workspace(tmp_path)
    ctx = resolve_context_sources(ws, "spec")
    pkg = build_document_phase_local_agent_package(
        phase="spec",
        payload={"runId": "r1", "idea_md": "# IDEA"},
        execution_policy={"requested": "local_agent_only", "selected": "local_agent", "reason": "t", "phase_supported": True},
    )
    out = attach_context_sources_to_local_package(pkg, ctx)
    local = out["local_agent"]
    files = {f["path"]: f for f in local["package_files"]}
    assert "PostgreSQL 16" in files["runs/context/spec/docs/inputs/Specifiche.docx.txt"]["content"]
    assert "runs/context/spec/docs/harper/ux/UI_DEMO.html.txt" not in files  # text files are read in place
    prompt = local["prompt_content"]
    assert "- runs/context/spec/docs/inputs/Specifiche.docx.txt — role: requirements (text converted from docs/inputs/Specifiche.docx)" in prompt
    assert "- docs/harper/ux/UI_DEMO.html — role: ui_reference" in prompt
    assert "- docs/harper/ux/assets/mock.png — role: ui_reference (image: open and look at it)" in prompt
    assert files[local["prompt_path"]]["content"] == prompt  # the prompt file matches
    assert {c["path"] for c in local["context_sources"]} >= {"docs/harper/ux/UI_DEMO.html"}
    assert attach_context_sources_to_local_package(pkg, {"items": []}) is pkg


def test_post_phase_run_adds_the_block_only_for_the_resolved_phase(monkeypatch, tmp_path):
    ctx = resolve_context_sources(_workspace(tmp_path), "spec")
    sent = {}

    async def fake_post(path, body):
        sent["body"] = body
        return {"ok": True}

    monkeypatch.setattr(harper, "_post_json", fake_post)
    monkeypatch.setattr(harper, "compose_phase_messages", lambda payload: [{"role": "user", "content": "go"}])

    asyncio.run(harper._post_phase_run({"phase": "spec", "context_sources": ctx}))
    assert "PROJECT CONTEXT SOURCES" in sent["body"]["composed_messages"][-1]["content"]
    assert "context_sources" not in sent["body"]

    asyncio.run(harper._post_phase_run({"phase": "integrity_eval", "context_sources": ctx}))
    assert sent["body"]["composed_messages"][-1]["content"] == "go"


def test_harper_init_template_manifest_is_valid():
    template = Path(__file__).resolve().parents[2] / "extensions/vscode/templates/harper-init"
    manifest = load_context_manifest(template)
    assert manifest["origin"] == ".clike/context.yaml"
    by_path = {s["path"]: s for s in manifest["sources"]}
    assert by_path["docs/inputs/**"]["role"] == "requirements"
    assert by_path["docs/harper/ux/**"] == {
        "path": "docs/harper/ux/**", "role": "ui_reference", "phases": ["spec", "plan", "kit"], "note": "",
    }
    assert manifest["limits"]["max_total_chars"] == 150_000
    # the template folders hold only .gitkeep, which is never read as a source
    assert resolve_context_sources(template, "spec")["items"] == []


def test_existing_source_code_is_read_as_context(tmp_path):
    # Brownfield modules and integration clients are context: a .py or .go file is not "unsupported".
    (tmp_path / ".clike").mkdir()
    (tmp_path / ".clike" / "context.yaml").write_text(
        "version: 1\nsources:\n  - path: legacy/**\n    role: existing_code\n    phases: [plan, kit]\n"
    )
    (tmp_path / "legacy").mkdir()
    (tmp_path / "legacy" / "billing.py").write_text("def invoice_number(year):\n    return f'{year}-0001'\n")
    (tmp_path / "legacy" / "client.go").write_text("package client\n")
    (tmp_path / "legacy" / "blob.bin").write_bytes(b"\x00\x01")
    context = resolve_context_sources(tmp_path, "kit")
    items = {i["path"]: i for i in context["items"]}
    assert items["legacy/billing.py"]["kind"] == "text"
    assert "invoice_number" in items["legacy/billing.py"]["text"]
    assert items["legacy/client.go"]["kind"] == "text"
    assert items["legacy/blob.bin"]["kind"] == "skipped"
    block = render_context_sources_for_cloud(context)
    assert "existing_code: code that already exists" in block
